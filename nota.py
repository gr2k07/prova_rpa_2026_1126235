
import tkinter as tk
from tkinter import ttk, messagebox
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# Todos os arquivos ficam na pasta deste programa.
PASTA = Path(__file__).resolve().parent
ARQUIVO = PASTA / "planilha_notas.xlsx"

COLUNAS = [
    "Nome", "RA", "Nota 1", "Nota 2",
    "Atividade 1", "Atividade 2", "Média", "Data do cadastro"
]


def calcular_media(nota1, nota2, atividade1, atividade2):
    return round((nota1 + nota2 + atividade1 + atividade2) / 4, 2)


def preparar_planilha():
    """Cria o arquivo Excel e seu cabeçalho, se necessário."""
    if not ARQUIVO.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Notas"
        ws.append(COLUNAS)

        for celula in ws[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill(
                "solid", fgColor="1D4ED8"
            )
            celula.alignment = Alignment(horizontal="center")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:H1"

        for coluna, largura in {
            "A": 28, "B": 16, "C": 12, "D": 12,
            "E": 16, "F": 16, "G": 12, "H": 22
        }.items():
            ws.column_dimensions[coluna].width = largura

        wb.save(ARQUIVO)
        wb.close()


def ler_alunos():
    """Lê os registros já salvos no Excel."""
    preparar_planilha()
    wb = load_workbook(ARQUIVO, data_only=True)
    ws = wb["Notas"]
    alunos = []

    for linha in ws.iter_rows(min_row=2, values_only=True):
        if linha[1] is not None:
            alunos.append(linha)

    wb.close()
    return alunos


def salvar_aluno(dados):
    """Grava um aluno na planilha."""
    preparar_planilha()
    wb = load_workbook(ARQUIVO)
    ws = wb["Notas"]

    # Impede duplicidade de RA.
    ra_novo = str(dados["ra"]).strip()
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if linha[1] is not None and str(linha[1]).strip() == ra_novo:
            wb.close()
            raise ValueError("Já existe um aluno cadastrado com esse RA.")

    media = calcular_media(
        dados["nota1"], dados["nota2"],
        dados["atividade1"], dados["atividade2"]
    )

    ws.append([
        dados["nome"],
        ra_novo,
        dados["nota1"],
        dados["nota2"],
        dados["atividade1"],
        dados["atividade2"],
        media,
        datetime.now().strftime("%d/%m/%Y %H:%M")
    ])

    wb.save(ARQUIVO)
    wb.close()


class SistemaNotas:
    def __init__(self, janela):
        self.janela = janela
        self.janela.title("Sistema Escolar - Cadastro de Notas")
        self.janela.geometry("1050x690")
        self.janela.minsize(780, 580)
        self.janela.configure(bg="#F1F5F9")

        self.campos = {}
        self.criar_interface()
        self.atualizar_tabela()

    def criar_interface(self):
        estilo = ttk.Style()
        estilo.theme_use("clam")
        estilo.configure(
            "Treeview",
            rowheight=29,
            font=("Arial", 10),
            background="white",
            fieldbackground="white"
        )
        estilo.configure(
            "Treeview.Heading",
            font=("Arial", 10, "bold"),
            background="#DBEAFE",
            foreground="#172033"
        )

        topo = tk.Frame(self.janela, bg="#1D4ED8", padx=22, pady=18)
        topo.pack(fill="x")

        tk.Label(
            topo,
            text="SISTEMA DE NOTAS ESCOLARES",
            bg="#1D4ED8", fg="white",
            font=("Arial", 19, "bold")
        ).pack(anchor="w")

        tk.Label(
            topo,
            text="Cadastro de alunos e controle de avaliações",
            bg="#1D4ED8", fg="#DBEAFE",
            font=("Arial", 10)
        ).pack(anchor="w", pady=(5, 0))

        formulario = tk.LabelFrame(
            self.janela,
            text=" Dados do aluno ",
            bg="white",
            fg="#172033",
            font=("Arial", 11, "bold"),
            padx=16, pady=14
        )
        formulario.pack(fill="x", padx=18, pady=14)

        nomes = [
            ("nome", "Nome completo"),
            ("ra", "RA"),
            ("nota1", "Nota 1 (0 a 10)"),
            ("nota2", "Nota 2 (0 a 10)"),
            ("atividade1", "Atividade 1 (0 a 10)"),
            ("atividade2", "Atividade 2 (0 a 10)")
        ]

        for indice, (chave, titulo) in enumerate(nomes):
            linha = indice // 3
            coluna = indice % 3

            area = tk.Frame(formulario, bg="white")
            area.grid(
                row=linha, column=coluna,
                sticky="ew", padx=8, pady=7
            )

            tk.Label(
                area, text=titulo,
                bg="white", fg="#334155",
                font=("Arial", 9, "bold")
            ).pack(anchor="w")

            entrada = ttk.Entry(area, font=("Arial", 11))
            entrada.pack(fill="x", pady=(5, 0), ipady=5)
            self.campos[chave] = entrada

        for coluna in range(3):
            formulario.columnconfigure(coluna, weight=1)

        botoes = tk.Frame(formulario, bg="white")
        botoes.grid(
            row=2, column=0, columnspan=3,
            sticky="w", padx=8, pady=(12, 2)
        )

        ttk.Button(
            botoes, text="Cadastrar aluno",
            command=self.cadastrar
        ).pack(side="left", padx=(0, 8))

        ttk.Button(
            botoes, text="Preencher exemplo",
            command=self.preencher_exemplo
        ).pack(side="left", padx=8)

        ttk.Button(
            botoes, text="Limpar campos",
            command=self.limpar
        ).pack(side="left", padx=8)

        area_tabela = tk.LabelFrame(
            self.janela,
            text=" Alunos registrados na planilha ",
            bg="white", fg="#172033",
            font=("Arial", 11, "bold"),
            padx=10, pady=10
        )
        area_tabela.pack(
            fill="both", expand=True, padx=18, pady=(0, 12)
        )

        nomes_colunas = (
            "nome", "ra", "nota1", "nota2",
            "ativ1", "ativ2", "media"
        )

        self.tabela = ttk.Treeview(
            area_tabela,
            columns=nomes_colunas,
            show="headings"
        )

        titulos = {
            "nome": "Nome",
            "ra": "RA",
            "nota1": "Nota 1",
            "nota2": "Nota 2",
            "ativ1": "Ativ. 1",
            "ativ2": "Ativ. 2",
            "media": "Média"
        }

        larguras = {
            "nome": 210, "ra": 100,
            "nota1": 75, "nota2": 75,
            "ativ1": 85, "ativ2": 85, "media": 80
        }

        for coluna in nomes_colunas:
            self.tabela.heading(coluna, text=titulos[coluna])
            self.tabela.column(
                coluna, width=larguras[coluna],
                minwidth=60, anchor="center"
            )

        self.tabela.column("nome", anchor="w")
        self.tabela.pack(
            side="left", fill="both", expand=True
        )

        barra = ttk.Scrollbar(
            area_tabela, orient="vertical",
            command=self.tabela.yview
        )
        barra.pack(side="right", fill="y")
        self.tabela.configure(yscrollcommand=barra.set)

        rodape = tk.Frame(self.janela, bg="#F1F5F9", padx=20)
        rodape.pack(fill="x", pady=(0, 12))

        self.status = tk.Label(
            rodape,
            text=f"Arquivo: {ARQUIVO.name}",
            bg="#F1F5F9", fg="#475569",
            font=("Arial", 9)
        )
        self.status.pack(side="left")

        ttk.Button(
            rodape, text="Atualizar tabela",
            command=self.atualizar_tabela
        ).pack(side="right")

    def obter_dados(self):
        dados = {
            chave: entrada.get().strip()
            for chave, entrada in self.campos.items()
        }

        if not dados["nome"] or not dados["ra"]:
            raise ValueError("Preencha o nome e o RA do aluno.")

        for chave in (
            "nota1", "nota2", "atividade1", "atividade2"
        ):
            try:
                valor = float(dados[chave].replace(",", "."))
            except ValueError:
                raise ValueError(
                    f"Digite um número válido em {chave}."
                )

            if not 0 <= valor <= 10:
                raise ValueError(
                    "Todas as notas e atividades devem ficar entre 0 e 10."
                )

            dados[chave] = valor

        return dados

    def cadastrar(self):
        try:
            dados = self.obter_dados()
            salvar_aluno(dados)
            self.atualizar_tabela()
            self.limpar()

            messagebox.showinfo(
                "Sucesso",
                f"Aluno {dados['nome']} cadastrado!\n\n"
                f"Planilha salva em:\n{ARQUIVO}"
            )

        except (ValueError, OSError, PermissionError) as erro:
            messagebox.showerror("Não foi possível cadastrar", str(erro))

    def preencher_exemplo(self):
        exemplo = {
            "nome": "Carlos Mendes",
            "ra": "2026004",
            "nota1": "8.5",
            "nota2": "9",
            "atividade1": "7.5",
            "atividade2": "10"
        }

        self.limpar()
        for chave, valor in exemplo.items():
            self.campos[chave].insert(0, valor)

    def limpar(self):
        for entrada in self.campos.values():
            entrada.delete(0, tk.END)

        self.campos["nome"].focus()

    def atualizar_tabela(self):
        for item in self.tabela.get_children():
            self.tabela.delete(item)

        try:
            alunos = ler_alunos()
            for aluno in alunos:
                nome, ra, n1, n2, a1, a2, media = aluno[:7]

                self.tabela.insert(
                    "", "end",
                    values=(
                        nome, ra,
                        f"{float(n1):.2f}",
                        f"{float(n2):.2f}",
                        f"{float(a1):.2f}",
                        f"{float(a2):.2f}",
                        f"{float(media):.2f}"
                    )
                )

            self.status.config(
                text=f"{len(alunos)} aluno(s) | Arquivo: {ARQUIVO.name}"
            )

        except (OSError, KeyError, ValueError) as erro:
            messagebox.showerror(
                "Erro ao ler planilha",
                f"Feche a planilha no Excel, se estiver aberta, "
                f"e tente novamente.\n\nDetalhes: {erro}"
            )


if __name__ == "__main__":
    preparar_planilha()
    janela = tk.Tk()
    app = SistemaNotas(janela)
    janela.mainloop()
